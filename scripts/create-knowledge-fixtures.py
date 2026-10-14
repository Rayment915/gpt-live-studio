"""Generate non-sensitive parser fixtures using ONLY the project venv."""
from pathlib import Path
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter

root = Path('output/knowledge-fixtures')
root.mkdir(parents=True, exist_ok=True)
text = '星河实验室：项目代号蓝鲸，预算7300元。设备借用期限为14天。'
for ext in ['txt', 'md', 'markdown']:
    (root / ('sample.' + ext)).write_text('# 非敏感测试资料\n' + text, encoding='utf-8')
(root / 'context-budget-test.md').write_text('# Non-sensitive context budget test\n' + 'This line is synthetic filler for transport testing only.\n' * 2400 + '\n结束标记：' + text + '\n', encoding='utf-8')
doc = Document()
doc.add_heading('星河实验室测试资料', 0)
doc.add_paragraph(text)
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = '项目'; table.cell(0, 1).text = '预算'
table.cell(1, 0).text = '蓝鲸'; table.cell(1, 1).text = '7300元'
doc.save(root / 'sample.docx')
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = '星河实验室测试资料'
slide.placeholders[1].text = text
prs.save(root / 'sample.pptx')
pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
c = canvas.Canvas(str(root / 'sample.pdf'))
c.setFont('STSong-Light', 18); c.drawString(60, 740, text[:25]); c.drawString(60, 705, text[25:]); c.showPage()
c.setFont('STSong-Light', 18); c.drawString(60, 740, '第二页：材料保管员为测试角色小林。'); c.save()
image = Image.new('RGB', (1400, 550), 'white'); draw = ImageDraw.Draw(image)
font = ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf', 40)
draw.text((50, 100), 'Scan only: Blue Whale budget = 7300 yuan.', font=font, fill='black')
draw.text((50, 200), 'Equipment loan period is 14 days.', font=font, fill='black')
image.save(root / 'scan.pdf', 'PDF', resolution=150)
encrypted = PdfWriter(); encrypted.append(PdfReader(root / 'sample.pdf')); encrypted.encrypt('fixture-password')
with (root / 'encrypted.pdf').open('wb') as target: encrypted.write(target)
for ext in ['pdf', 'docx', 'pptx', 'doc', 'ppt']:
    (root / ('corrupt.' + ext)).write_bytes(b'not a valid office or pdf file')
(root / 'oversized.txt').write_bytes(b'x' * (20 * 1024 * 1024 + 1))
print('Non-sensitive fixtures saved in', root)
